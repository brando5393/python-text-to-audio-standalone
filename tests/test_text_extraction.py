import os

import docx
import pypdf
import pytest
from ebooklib import epub
from odf.opendocument import OpenDocumentText
from odf.text import P as OdfParagraph
from pptx import Presentation

import TextExtraction as te


def test_extract_txt(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Plain text content.", encoding="utf-8")
    assert te.extract_text(str(path)) == "Plain text content."


def test_extract_markdown_strips_common_syntax(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("# Heading\n\nSome **bold** and _italic_ text.\n", encoding="utf-8")
    result = te.extract_text(str(path))
    assert "#" not in result
    assert "*" not in result
    assert "Heading" in result
    assert "bold" in result


def test_extract_markdown_converts_links_to_plain_text(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("See [the docs](https://example.com) for more.", encoding="utf-8")
    result = te.extract_text(str(path))
    assert "https://example.com" not in result
    assert "the docs" in result


def test_extract_docx(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    result = te.extract_text(str(path))
    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_extract_rtf(tmp_path):
    path = tmp_path / "note.rtf"
    path.write_text(r"{\rtf1\ansi Hello from RTF.\par}", encoding="utf-8")
    result = te.extract_text(str(path))
    assert "Hello from RTF." in result


def test_extract_html(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<html><body><h1>Title</h1><p>Body text.</p></body></html>", encoding="utf-8")
    result = te.extract_text(str(path))
    assert "Title" in result
    assert "Body text." in result


def test_extract_epub(tmp_path):
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Test Book")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter 1", file_name="chap1.xhtml", lang="en")
    chapter.content = "<h1>Chapter One</h1><p>Some epub content.</p>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    path = tmp_path / "book.epub"
    epub.write_epub(str(path), book)
    result = te.extract_text(str(path))
    assert "Some epub content." in result


def test_extract_xhtml(tmp_path):
    path = tmp_path / "page.xhtml"
    path.write_text(
        '<?xml version="1.0"?><html><body><h1>Title</h1><p>Body text.</p></body></html>',
        encoding="utf-8",
    )
    result = te.extract_text(str(path))
    assert "Title" in result
    assert "Body text." in result


_FB2_TEMPLATE = """<?xml version="1.0" encoding="UTF-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">
  <body>
    <section><title><p>Chapter One</p></title><p>First chapter prose.</p></section>
    <section><title><p>Chapter Two</p></title><p>Second chapter prose.</p></section>
  </body>
  <body name="notes">
    <section><p>A footnote, not part of the main text.</p></section>
  </body>
</FictionBook>
"""


def _make_fb2(path):
    path.write_text(_FB2_TEMPLATE, encoding="utf-8")


def test_extract_fb2(tmp_path):
    path = tmp_path / "book.fb2"
    _make_fb2(path)
    result = te.extract_text(str(path))
    assert "First chapter prose." in result
    assert "Second chapter prose." in result
    assert "A footnote" not in result


def test_extract_pptx(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide Title"
    slide.placeholders[1].text_frame.text = "Body text"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    result = te.extract_text(str(path))
    assert "Slide Title" in result
    assert "Body text" in result


def test_extract_odt(tmp_path):
    doc = OpenDocumentText()
    doc.text.addElement(OdfParagraph(text="First paragraph."))
    doc.text.addElement(OdfParagraph(text="Second paragraph."))
    path = tmp_path / "doc.odt"
    doc.save(str(path))
    result = te.extract_text(str(path))
    assert "First paragraph." in result
    assert "Second paragraph." in result


@pytest.mark.parametrize("ext", [".azw", ".prc"])
def test_azw_and_prc_route_through_the_mobi_extractor(tmp_path, monkeypatch, ext):
    # .azw and .prc are the same underlying Palm-database Kindle format as .mobi/.azw3 --
    # the `mobi` library sniffs the file's binary header, not its extension -- so real
    # extraction is already covered by the .mobi/.azw3 code path; this just confirms
    # dispatch reaches it instead of raising ValueError.
    calls = []
    monkeypatch.setattr(te, "_extract_mobi", lambda path: calls.append(path) or "extracted")
    path = tmp_path / f"book{ext}"
    path.write_bytes(b"fake mobi bytes")
    assert te.extract_text(str(path)) == "extracted"
    assert calls == [str(path)]


def test_unsupported_extension_raises_value_error(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError):
        te.extract_text(str(path))


def test_supported_extensions_are_all_lowercase_with_dot():
    for ext in te.SUPPORTED_EXTENSIONS:
        assert ext.startswith(".")
        assert ext == ext.lower()


def _make_pdf(path, num_pages):
    writer = pypdf.PdfWriter()
    for _ in range(num_pages):
        writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)


def _make_epub(path, num_chapters):
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Test Book")
    book.set_language("en")
    chapters = []
    for i in range(num_chapters):
        chapter = epub.EpubHtml(title=f"Chapter {i + 1}", file_name=f"chap{i + 1}.xhtml", lang="en")
        chapter.content = f"<h1>Chapter {i + 1}</h1><p>Content for chapter {i + 1}.</p>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + chapters
    epub.write_epub(str(path), book)


def test_extract_structure_counts_pdf_returns_page_count(tmp_path):
    path = tmp_path / "doc.pdf"
    _make_pdf(path, 5)
    assert te.extract_structure_counts(str(path)) == (5, None)


def test_extract_structure_counts_epub_returns_chapter_count(tmp_path):
    path = tmp_path / "book.epub"
    _make_epub(path, 4)
    assert te.extract_structure_counts(str(path)) == (None, 4)


def test_extract_structure_counts_docx_counts_heading_1_paragraphs(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("Some prose here.")
    doc.add_paragraph("Chapter Two", style="Heading 1")
    doc.add_paragraph("More prose here.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    assert te.extract_structure_counts(str(path)) == (None, 2)


def test_extract_structure_counts_docx_returns_none_without_headings(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Just a plain paragraph with no heading style.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    assert te.extract_structure_counts(str(path)) == (None, None)


def test_extract_structure_counts_fb2_returns_chapter_count(tmp_path):
    path = tmp_path / "book.fb2"
    _make_fb2(path)
    assert te.extract_structure_counts(str(path)) == (None, 2)


def test_extract_structure_counts_pptx_returns_slide_count(tmp_path):
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    assert te.extract_structure_counts(str(path)) == (3, None)


def test_extract_structure_counts_txt_returns_none_none(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Plain text content.", encoding="utf-8")
    assert te.extract_structure_counts(str(path)) == (None, None)


def test_estimate_structure_returns_none_none_for_empty_text():
    assert te.estimate_structure("") == (None, None)
    assert te.estimate_structure("   \n  ") == (None, None)


def test_estimate_structure_estimates_pages_from_character_count():
    pages, chapters = te.estimate_structure("x" * 9000)
    assert pages == 3  # 9000 chars / 3000 chars-per-page
    assert chapters is None  # no chapter-heading-looking lines present


def test_estimate_structure_never_returns_zero_pages_for_short_text():
    pages, _ = te.estimate_structure("A single short sentence.")
    assert pages == 1


def test_estimate_structure_detects_chapter_heading_lines():
    text = "Chapter 1\nSome text.\n\nCHAPTER TWO\nMore text.\n\nChapter III:\nEven more."
    pages, chapters = te.estimate_structure(text)
    assert chapters == 3


def test_estimate_structure_ignores_a_single_chapter_heading_line():
    # One matching line alone is too weak a signal to trust (could be a one-off title
    # page rather than real recurring chapter structure) -- needs at least two.
    text = "Chapter 1\n" + ("Filler text. " * 50)
    _, chapters = te.estimate_structure(text)
    assert chapters is None


def test_estimate_structure_does_not_match_chapter_mentioned_mid_sentence():
    text = "This chapter covers a lot of ground.\n" + ("Filler text. " * 50)
    _, chapters = te.estimate_structure(text)
    assert chapters is None


def test_extract_chapters_epub_returns_one_text_per_chapter(tmp_path):
    path = tmp_path / "book.epub"
    _make_epub(path, 3)
    chapters = te.extract_chapters(str(path))
    assert len(chapters) == 3
    assert "Content for chapter 1." in chapters[0]
    assert "Content for chapter 2." in chapters[1]
    assert "Content for chapter 3." in chapters[2]


def test_extract_chapters_docx_splits_on_heading_1(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("First chapter prose.")
    doc.add_paragraph("Chapter Two", style="Heading 1")
    doc.add_paragraph("Second chapter prose.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    chapters = te.extract_chapters(str(path))
    assert len(chapters) == 2
    assert "First chapter prose." in chapters[0]
    assert "Second chapter prose." not in chapters[0]
    assert "Second chapter prose." in chapters[1]


def test_extract_chapters_docx_keeps_text_before_first_heading(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Front matter before any heading.")
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("First chapter prose.")
    doc.add_paragraph("Chapter Two", style="Heading 1")
    doc.add_paragraph("Second chapter prose.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    chapters = te.extract_chapters(str(path))
    assert len(chapters) == 3
    assert "Front matter before any heading." in chapters[0]


def test_extract_chapters_docx_returns_none_without_headings(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Just a plain paragraph with no heading style.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    assert te.extract_chapters(str(path)) is None


def test_extract_chapters_fb2_returns_one_text_per_chapter(tmp_path):
    path = tmp_path / "book.fb2"
    _make_fb2(path)
    chapters = te.extract_chapters(str(path))
    assert len(chapters) == 2
    assert "First chapter prose." in chapters[0]
    assert "Second chapter prose." not in chapters[0]
    assert "Second chapter prose." in chapters[1]


def test_extract_chapters_returns_none_for_unsupported_format(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Plain text content.", encoding="utf-8")
    assert te.extract_chapters(str(path)) is None


def test_extract_chapters_pdf_returns_none(tmp_path):
    path = tmp_path / "doc.pdf"
    _make_pdf(path, 3)
    assert te.extract_chapters(str(path)) is None
