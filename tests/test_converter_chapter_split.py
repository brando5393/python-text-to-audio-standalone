"""Verifies split_chapters=True converts an eligible EPUB/DOCX file into one audio file
per chapter, and that ineligible formats (or documents with no detectable chapter
structure) fall back to a normal single whole-file conversion instead of erroring."""

import json
import time
import wave

import docx
from ebooklib import epub

import Config
import Converter
import PiperEngine


def _write_fake_wav(_text, _voice_id, path, **_kwargs):
    with wave.open(path, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(22050)
        w.writeframes(b"\x00\x00" * 100)


def _wait_for_all_done(converter, timeout=20):
    events = []
    deadline = time.time() + timeout
    while time.time() < deadline:
        new_events = converter.poll_events()
        events.extend(new_events)
        if any(e[0] == "all_done" for e in new_events):
            return events
        time.sleep(0.05)
    raise TimeoutError(f"Conversion did not finish within {timeout}s. Events so far: {events}")


def _make_converter(tk_root, tmp_path, monkeypatch, **config_overrides):
    monkeypatch.setattr(Config, "CONFIG_PATH", str(tmp_path / "config.json"))
    Config.save({**Config.DEFAULTS, **config_overrides})
    monkeypatch.setattr(PiperEngine, "is_engine_installed", lambda: True)
    monkeypatch.setattr(PiperEngine, "is_voice_installed", lambda voice: True)
    monkeypatch.setattr(PiperEngine, "synthesize", _write_fake_wav)

    import tkinter as tk
    return Converter.Converter(tk.Listbox(tk_root))


def _make_epub(path, num_chapters):
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Test Book")
    book.set_language("en")
    chapters = []
    for i in range(num_chapters):
        chapter = epub.EpubHtml(title=f"Chapter {i + 1}", file_name=f"chap{i + 1}.xhtml", lang="en")
        chapter.content = f"<h1>Chapter {i + 1}</h1><p>Content for chapter {i + 1}.</p>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + chapters
    epub.write_epub(str(path), book)


def test_split_chapters_epub_produces_one_file_per_chapter(tk_root, tmp_path, monkeypatch):
    converter = _make_converter(tk_root, tmp_path, monkeypatch, engine="piper", voice="en_US-amy-medium")

    src = tmp_path / "book.epub"
    _make_epub(src, 3)
    out_dir = tmp_path / "out"

    converter.convert_to_audio([str(src)], str(out_dir), split_chapters=True)
    events = _wait_for_all_done(converter)

    assert "error" not in [e[0] for e in events]
    assert (out_dir / "book - Chapter 1.wav").exists()
    assert (out_dir / "book - Chapter 2.wav").exists()
    assert (out_dir / "book - Chapter 3.wav").exists()
    assert not (out_dir / "book.wav").exists()

    with open(out_dir / "book - Chapter 1.wav.json", encoding="utf-8") as f:
        sidecar = json.load(f)
    assert sidecar["engine"] == "piper"
    assert "Content for chapter 1." in sidecar["text"]


def test_split_chapters_docx_produces_one_file_per_chapter(tk_root, tmp_path, monkeypatch):
    converter = _make_converter(tk_root, tmp_path, monkeypatch, engine="piper", voice="en_US-amy-medium")

    doc = docx.Document()
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("First chapter prose.")
    doc.add_paragraph("Chapter Two", style="Heading 1")
    doc.add_paragraph("Second chapter prose.")
    src = tmp_path / "doc.docx"
    doc.save(str(src))
    out_dir = tmp_path / "out"

    converter.convert_to_audio([str(src)], str(out_dir), split_chapters=True)
    _wait_for_all_done(converter)

    assert (out_dir / "doc - Chapter 1.wav").exists()
    assert (out_dir / "doc - Chapter 2.wav").exists()
    assert not (out_dir / "doc.wav").exists()


def test_split_chapters_falls_back_to_whole_file_for_ineligible_format(tk_root, tmp_path, monkeypatch):
    """A .txt file has no chapter concept at all -- split_chapters=True must not error,
    it should just convert the whole file normally, same as split_chapters=False."""
    converter = _make_converter(tk_root, tmp_path, monkeypatch, engine="piper", voice="en_US-amy-medium")

    src = tmp_path / "note.txt"
    src.write_text("Some plain text with no chapter structure at all.", encoding="utf-8")
    out_dir = tmp_path / "out"

    converter.convert_to_audio([str(src)], str(out_dir), split_chapters=True)
    events = _wait_for_all_done(converter)

    assert "error" not in [e[0] for e in events]
    assert (out_dir / "note.wav").exists()


def test_split_chapters_falls_back_for_docx_without_headings(tk_root, tmp_path, monkeypatch):
    """A .docx with no Heading 1 paragraphs has no detectable chapter boundary -- this
    should convert as one whole file instead of failing or producing zero output."""
    converter = _make_converter(tk_root, tmp_path, monkeypatch, engine="piper", voice="en_US-amy-medium")

    doc = docx.Document()
    doc.add_paragraph("Just a plain paragraph with no heading style.")
    src = tmp_path / "doc.docx"
    doc.save(str(src))
    out_dir = tmp_path / "out"

    converter.convert_to_audio([str(src)], str(out_dir), split_chapters=True)
    events = _wait_for_all_done(converter)

    assert "error" not in [e[0] for e in events]
    assert (out_dir / "doc.wav").exists()


def test_default_split_chapters_is_false_and_behaves_as_before(tk_root, tmp_path, monkeypatch):
    converter = _make_converter(tk_root, tmp_path, monkeypatch, engine="piper", voice="en_US-amy-medium")

    src = tmp_path / "book.epub"
    _make_epub(src, 2)
    out_dir = tmp_path / "out"

    converter.convert_to_audio([str(src)], str(out_dir))  # split_chapters not passed
    events = _wait_for_all_done(converter)

    assert "error" not in [e[0] for e in events]
    assert (out_dir / "book.wav").exists()
    assert not (out_dir / "book - Chapter 1.wav").exists()
